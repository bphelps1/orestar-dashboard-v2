#!/usr/bin/env python3
"""
match_lobbyists.py — suggest which lobbyist handles which ORESTAR donor.

Reads lobbyists / lobbyist_clients (fetch_capitol_club.py), committee_persons
(fetch_committee_persons.py) and the donor pool, and writes suggestions to
donor_lobbyist_links and donor_client_links for an admin to confirm at
/admin/lobbyists. A row a human has confirmed or rejected is never touched;
a machine suggestion that the current data no longer supports is removed.

Evidence, strongest first:

  committee contact → lobbyist   (donor committees only)
    email_exact   a treasurer / correspondent / director email equals the
                  lobbyist's Capitol Club email        (Sean Kolmer → OHPAC 161)
    name_exact    same first + last name
    email_domain  same private email domain            (freelandern@seiu503.org
                  → Courtney Graham, grahamc@seiu503.org). Mail providers,
                  .gov/.us/.edu and treasurer-service domains never count.
    director      a director's employer is one of the lobbyist's clients

  donor name → Capitol Club client
    name_exact    identical after dropping legal suffixes, PAC wording and
                  committee ids ("Kroger" / "The Kroger Co.")
    name_fuzzy    IDF-weighted token overlap in both directions (see
                  fuzzy_score), sharing the client's rarest word and adding no
                  place the client lacks. Public-body clients never match.

  seeds (local files, never committed — the repo is public)
    --tracker   Fundraising Tracker .xlsx, "Lobbyist Key" F:I
                (contributor, committee id, lobbyist 1, lobbyist 2) — curated
                by the fundraising team, so stored as confirmed.
    --sheet2024 JSON export of the 2024 FuturePAC lobby list — adds the
                lobbyists Capitol Club lacks (e.g. union partners) and their
                clients; for lobbyists Capitol Club does list, the 2024
                client pairs are kept only as inactive history.

Usage:
    python scraper/match_lobbyists.py --refresh-pool
    python scraper/match_lobbyists.py --tracker ~/tracker.xlsx --sheet2024 lobby2024.json
    python scraper/match_lobbyists.py --dry-run --report out.json
"""

from __future__ import annotations

import argparse
import json
import logging
import math
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
import supabase_sync  # noqa: E402
from lobby_match import (  # noqa: E402
    client_alternatives, core_org, dba_names, email_domain, first_last, first_names_compatible,
    is_private_domain, is_public_client, names_same_org, nicknames, norm_org, norm_person,
    person_label, person_tokens, place_tokens,
)

log = logging.getLogger(__name__)

POOL_SINCE = "2021-01-01"
FUZZY_MIN_F = 0.8         # balanced (both-direction) IDF-weighted overlap
FUZZY_ANCHORED = 0.35     # …or: same leading word, all of the client, this much of the donor
FUZZY_MIN_IDF = 7.0       # total IDF the shared tokens must carry
BRAND_MIN_IDF = 7.5       # a one-word donor name this rare can stand for its brand
PERSON_ROLES = ("treasurer", "correspondence", "director")
TRACKER_DECIDER = "Fundraising Tracker — Lobbyist Key"

REFRESH_POOL_SQL = f"""
with t as (
  select donor_id, amount, filer_id, tran_date,
         lower(regexp_replace(btrim(coalesce(nullif(contributor_payee_canonical, ''),
                                                contributor_payee)), '\\s+', ' ', 'g')) as label,
         addr_line1
  from transactions
  where tran_date >= '{POOL_SINCE}'
    and sub_type in ('Cash Contribution', 'In-Kind Contribution')
    and coalesce(book_type, '') not in ('Individual', 'Candidate & Immediate Family')
    and donor_id is not null
    and not (contributor_payee ilike 'miscellaneous %%')
)
insert into lobby_donor_pool (donor_id, display_name, book_type, committee_id, names, address,
                              city, state, total_since_2021, gifts, recipients, last_date, refreshed_at)
select d.donor_id, d.display_name, d.book_type, nullif(d.committee_id, ''),
       array_agg(distinct t.label), mode() within group (order by t.addr_line1),
       d.city, d.state, sum(t.amount), count(*), count(distinct t.filer_id), max(t.tran_date), now()
from t join donors d using (donor_id)
group by d.donor_id, d.display_name, d.book_type, d.committee_id, d.city, d.state
"""


# ── Loading ──────────────────────────────────────────────────────────────────

def _rows(cur, sql, params=()):
    cur.execute(sql, params)
    cols = [c[0] for c in cur.description]
    return [dict(zip(cols, r)) for r in cur.fetchall()]


def refresh_pool(conn) -> int:
    cur = conn.cursor()
    cur.execute("set statement_timeout = 0")
    cur.execute("delete from lobby_donor_pool")
    cur.execute(REFRESH_POOL_SQL)
    n = cur.rowcount
    conn.commit()
    return n


# ── Name index for donor ↔ client matching ───────────────────────────────────

class OrgIndex:
    """Token IDF over donor and client names, for weighted overlap."""

    def __init__(self, names: list[str]):
        df = Counter()
        for n in names:
            df.update(set(n.split()))
        self.n = max(len(names), 1)
        self.df = df

    def idf(self, tok: str) -> float:
        return math.log((self.n + 1) / (self.df.get(tok, 0) + 1)) + 1.0

    def overlap(self, donor: str, client: str) -> tuple[float, float, float]:
        """(share of the client covered, share of the donor covered, shared IDF)."""
        td, tc = set(donor.split()), set(client.split())
        common = td & tc
        if not common:
            return 0.0, 0.0, 0.0
        w = lambda s: sum(self.idf(t) for t in s)
        shared = w(common)
        return shared / w(tc), shared / w(td), shared


def fuzzy_score(idx: OrgIndex, donor: str, client: str) -> float:
    """0 when not a plausible match, else a 0.5–0.9 suggestion score.

    Containment of the shorter name alone is not enough: "Apple" sits wholly
    inside "Apple City Auto Body", "Age" inside "Space Age Fuel". Either both
    names must be mostly covered (balanced F), or they must start with the
    same word, the client must be fully covered and the donor's extra words
    must be modest ("Everytown for Gun Safety" ↔ "Everytown").
    """
    c_client, c_donor, shared = idx.overlap(donor, client)
    if shared < FUZZY_MIN_IDF and not (c_client == 1.0 and c_donor == 1.0):
        return 0.0
    # A place the client lacks makes the donor a local outfit or affiliate.
    if place_tokens(donor) - place_tokens(client):
        return 0.0
    anchored = donor.split()[0] == client.split()[0]
    toks = donor.split()
    top_client_idf = max(idx.idf(t) for t in client.split())
    # A one-word brand that leads the client's name and is its most
    # distinctive word: "Davita" ↔ "DaVita HealthCare Partners", "Regence" ↔
    # "Regence BlueCross BlueShield" — but not "Williams" ↔ "Williams &
    # Russell CDC", whose rarer words say it is someone else.
    if (anchored and len(toks) == 1 and idx.idf(toks[0]) >= BRAND_MIN_IDF
            and idx.idf(toks[0]) >= top_client_idf - 0.5):
        return round(0.5 + 0.3 * c_client, 3)
    # Otherwise the client's most distinctive word must be shared:
    # "Associated Oregon Industries" is not "Associated Oregon Hazelnut
    # Industries".
    if max(client.split(), key=idx.idf) not in toks:
        return 0.0
    f = 2 * c_client * c_donor / (c_client + c_donor) if c_client + c_donor else 0.0
    if f >= FUZZY_MIN_F:
        return round(0.5 + 0.4 * f, 3)
    if anchored and c_client >= 0.99 and c_donor >= FUZZY_ANCHORED:
        return round(0.5 + 0.3 * c_donor, 3)
    return 0.0


def match_clients(pool: list[dict], clients: dict[str, str]) -> list[dict]:
    """Donor ↔ client name suggestions. clients: client_key → display name."""
    # client_key → comparison cores (a "/"-joined client yields several)
    client_cores: dict[str, set[str]] = {}
    for k, v in clients.items():
        if is_public_client(v):
            continue
        cores = {core_org(a) for a in client_alternatives(v)} - {""}
        if cores:
            client_cores[k] = cores
    donor_cores: dict[str, set[str]] = {}
    for d in pool:
        cores = set()
        for nm in [d["display_name"], *(d["names"] or [])]:
            for part in dba_names(nm):
                c = core_org(part)
                if c:
                    cores.add(c)
        donor_cores[d["donor_id"]] = cores

    idx = OrgIndex([c for cs in client_cores.values() for c in cs]
                   + [c for cs in donor_cores.values() for c in cs])
    by_token: dict[str, set[str]] = defaultdict(set)
    exact: dict[str, set[str]] = defaultdict(set)
    for k, cores in client_cores.items():
        for c in cores:
            exact[c].add(k)
            for t in c.split():
                by_token[t].add(k)

    out = []
    for d in pool:
        best: dict[str, tuple] = {}
        for core in donor_cores[d["donor_id"]]:
            for k in exact.get(core, ()):
                best[k] = ("name_exact", 0.95, core, core)
            # Candidates must share one of the donor's two rarest tokens;
            # common words ("oregon", "association") are too broad to block on.
            toks = sorted(core.split(), key=idx.idf, reverse=True)[:2]
            for t in toks:
                if idx.idf(t) < 4.0:
                    continue
                for k in by_token.get(t, ()):
                    if k in best and best[k][0] == "name_exact":
                        continue
                    for ccore in client_cores[k]:
                        s = fuzzy_score(idx, core, ccore)
                        if s and (k not in best or best[k][1] < s):
                            best[k] = ("name_fuzzy", s, core, ccore)
        for k, (method, score, dcore, ccore) in best.items():
            out.append({"donor_id": d["donor_id"], "client_key": k, "client_name": clients[k],
                        "method": method, "score": score,
                        "evidence": [{"type": method, "donor_name": d["display_name"],
                                      "donor_core": dcore, "client_core": ccore}]})
    return out


# ── Committee contacts ↔ lobbyists ───────────────────────────────────────────

def match_committee_contacts(pool, persons, lobbyists, lobbyist_clients) -> tuple[list, list]:
    by_committee = {d["committee_id"]: d for d in pool if d["committee_id"]}
    by_email = defaultdict(list)
    by_name = defaultdict(list)
    by_domain = defaultdict(list)
    for l in lobbyists:
        if l["kind"] != "person":
            continue
        if l["email"]:
            by_email[l["email"].lower()].append(l)
            dom = email_domain(l["email"])
            if is_private_domain(dom):
                by_domain[dom].append(l)
        f, la = first_last(l["name"])
        if f and la:
            by_name[(f, la)].append(l)
    client_to_lobbyists = defaultdict(set)
    client_names = {}
    for c in lobbyist_clients:
        if c["active"]:
            client_to_lobbyists[core_org(c["client_name"])].add(c["lobbyist_id"])
            client_names[core_org(c["client_name"])] = (c["client_key"], c["client_name"])

    links, client_links = [], []
    for p in persons:
        d = by_committee.get(p["filer_id"])
        if not d or p["role"] not in PERSON_ROLES:
            continue
        who = {"role": p["role"], "person": p["name"], "email": p["email"] or ""}
        hits: dict[int, tuple] = {}
        email = (p["email"] or "").lower()
        for l in by_email.get(email, []) if email else []:
            hits[l["lobbyist_id"]] = ("email_exact", 0.98)
        f, la = first_last(p["name"])
        for l in by_name.get((f, la), []):
            if l["lobbyist_id"] not in hits:
                hits[l["lobbyist_id"]] = ("name_exact", 0.9 if p["role"] != "director" else 0.8)
        dom = email_domain(email)
        if is_private_domain(dom):
            peers = by_domain.get(dom, [])
            # A domain shared by a dozen lobbyists is a big firm or agency;
            # still a lead, but weaker.
            s = 0.7 if len(peers) <= 3 else 0.55
            for l in peers:
                if l["lobbyist_id"] not in hits:
                    hits[l["lobbyist_id"]] = ("email_domain", s)
        for lid, (method, score) in hits.items():
            links.append({"donor_id": d["donor_id"], "lobbyist_id": lid, "method": method,
                          "score": score, "evidence": [{"type": method, **who}]})
        # A director's employer is the committee's client only when the
        # committee is that organization's own PAC; a chamber's or trade
        # group's board is made of people from other companies.
        if p["role"] == "director" and p["employer"] and not is_public_client(p["employer"]) \
                and names_same_org(d["display_name"], p["employer"]):
            core = core_org(p["employer"])
            if core in client_names:
                key, cname = client_names[core]
                client_links.append({"donor_id": d["donor_id"], "client_key": key, "client_name": cname,
                                     "method": "committee_contact", "score": 0.75,
                                     "evidence": [{"type": "director_employer", **who,
                                                   "employer": p["employer"]}]})
    return links, client_links


# ── Seeds ────────────────────────────────────────────────────────────────────

_FIRM_WORDS = re.compile(r"\b(associates|affairs|group|strategies|torp|west|counsel|cfm|run|"
                         r"communications|communcations|partners|relations|consulting|lobby)\b|&", re.I)


def match_person(label: str, lobbyists: list[dict], prefer_emails: set[str] = frozenset()) -> list[dict]:
    """The lobbyist(s) a name refers to, or [] when unknown or ambiguous.

    Handles "LAST, FIRST" and "First Last", compound surnames ("PALMATEER,
    NICOLE" → Nicole Palmateer-Hazelbaker), spacing ("Desitter" → "De
    Sitter"), one-letter typos ("DEMSEY"), nicknames ("Selvaggio, Mike" →
    Michael Selvaggio) and, when the surname is unique, a bare first initial
    ("LOVING, DAN" → Don Loving). Several rows that normalize to one person
    (Capitol Club lists Eliza Walton twice) all match.
    """
    from rapidfuzz.fuzz import ratio
    first, lasts = person_label(label)
    if not lasts:
        return []
    persons = [l for l in lobbyists if l["kind"] == "person"]

    def surname_hit(l, fuzzy=False):
        toks = person_tokens(l["name"])
        joined = "".join(toks[1:])
        for last in lasts:
            if last in toks[1:] or last == joined or "".join(lasts) == joined:
                return True
            if fuzzy and len(last) >= 5 and any(ratio(last, t) >= 85 for t in toks[1:]):
                return True
        return False

    def first_of(l):
        toks = person_tokens(l["name"])
        return toks[0] if toks else ""

    def first_ok(l):
        return first_names_compatible(first, first_of(l)) or first in nicknames(l["name"])

    if not first:   # single-word label: a surname, or a first name / nickname alone
        cands = [l for l in persons if surname_hit(l)]
        cands = cands or [l for l in persons if first_of(l) == lasts[0] or lasts[0] in nicknames(l["name"])]
    else:
        cands = [l for l in persons if surname_hit(l) and first_ok(l)]
        if not cands:
            cands = [l for l in persons if surname_hit(l, fuzzy=True) and first_ok(l)]
        if not cands:
            same_surname = [l for l in persons if surname_hit(l)]
            if len({norm_person(l["name"]) for l in same_surname}) == 1 \
                    and first_of(same_surname[0])[:1] == first[:1]:
                cands = same_surname
    groups: dict[str, list[dict]] = defaultdict(list)
    for l in cands:
        groups[norm_person(l["name"])].append(l)
    if len(groups) > 1:
        on_cc = {k: v for k, v in groups.items() if any(l["on_capitol_club"] for l in v)}
        groups = on_cc or groups
    if len(groups) > 1 and prefer_emails:
        pref = {k: v for k, v in groups.items() if any((l["email"] or "").lower() in prefer_emails for l in v)}
        groups = pref or groups
    return next(iter(groups.values())) if len(groups) == 1 else []


def resolve_label(label: str, lobbyists: list[dict],
                  prefer_emails: set[str] = frozenset()) -> tuple[str, list[dict]]:
    """Tracker label → ('person'|'firm', matching lobbyists)."""
    raw = re.sub(r"\s+S\d\b.*$", "", label.strip())      # "ANGSTROM, RICH S1"
    if _FIRM_WORDS.search(raw):
        key = norm_org(raw)
        return "firm", [l for l in lobbyists if l["kind"] == "firm" and (
            norm_org(l["name"]) == key or key in {norm_org(a) for a in l["aliases"] or []})]
    return "person", match_person(raw, lobbyists, prefer_emails)


def load_tracker(path: Path) -> list[tuple[str, str, str, str]]:
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb["Lobbyist Key"]
    out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        r = list(r) + [None] * 9           # read-only mode trims empty trailing cells
        contributor, cid, l1, l2 = r[5], r[6], r[7], r[8]
        if not contributor:
            continue
        labels = [x for x in (l1, l2) if x and str(x).strip() not in ("-", "")]
        if labels:
            out.append((str(contributor).strip(), str(cid or "").strip(), *[str(x) for x in labels]))
    return out


def load_tracker_notes(path: Path) -> list[str]:
    """Lobbyist Key column A: "THORN RUN - Dan Bates S3", "OXLEY & ASSOCIATES (Evyan) S2"."""
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    notes = []
    for r in wb["Lobbyist Key"].iter_rows(min_row=2, values_only=True):
        if r and r[0] and str(r[0]).strip() not in notes:
            notes.append(str(r[0]).strip())
    return notes


# ── Firm contacts and client leads ───────────────────────────────────────────
#
# A plan files some donors under a firm ("THORN RUN") and needs a person to
# call there; a client listed by several lobbyists (the Hospital Association
# lists three) needs one to file its donors under. The fundraising sheets
# already make both choices, so they seed them. Both are only filled where an
# admin has not set them, and only with people Capitol Club still places at
# the firm: the 2024 list's Tonkon Torp lead has since moved to another firm.

_PHONE = re.compile(r"\(?\d{3}\)?[\s.-]*\d{3}[\s.-]\d{4}")
_FIRM_STOP = {"and", "associates", "public", "affairs", "group", "the", "strategies", "government",
              "relations", "communications", "partners", "lobby", "counsel", "consulting", "inc",
              "llc", "co", "company", "strategic", "advocates"}


def _strip_tag(label: str) -> str:
    """Drop the tracker's tier tag and phone numbers: "CFM - Dale Penn 503-510-2200 S2"."""
    label = _PHONE.sub(" ", label)
    return re.sub(r"\s+S\d\b.*$|\s+-\s*$", "", label.strip()).strip(" -")


def _firm_phrase(name: str) -> str:
    return norm_org(name.replace("Communcations", "Communications").replace("COMMUNCATIONS", "COMMUNICATIONS"))


def parse_firm_note(note: str) -> tuple[str, list[str]]:
    """"OXLEY & ASSOCIATES (Evyan) S2" → ("OXLEY & ASSOCIATES", ["Evyan"])."""
    raw = _strip_tag(note)
    m = re.match(r"^(.*?)\s*\(([^)]*)\)", raw)
    if m:
        firm, people = m.group(1), m.group(2)
    elif " - " in raw:
        firm, people = raw.split(" - ", 1)
    else:
        return raw, []
    return firm.strip(), [p.strip() for p in re.split(r"[/;]", _PHONE.sub(" ", people)) if p.strip()]


def _names_in(text: str) -> list[str]:
    """Person names in a free-text "Additional Lobbyists" cell."""
    out = []
    for chunk in re.split(r"[;,:]|" + _PHONE.pattern, text or ""):
        toks = [t for t in re.split(r"\s+", chunk.strip(" -.")) if t]
        i = 0
        while i + 1 < len(toks):
            out.append(f"{toks[i]} {toks[i + 1]}")
            i += 2
    return out


def seed_firm_contacts(cur, tracker_rows, notes: list[str], sheet2024: list[dict]) -> None:
    lobbyists = _rows(cur, "select * from lobbyists")
    persons = [l for l in lobbyists if l["kind"] == "person"]
    by_id = {l["lobbyist_id"]: l for l in persons}
    n_clients = Counter(r[0] for r in _rows_raw(cur, "select lobbyist_id from lobbyist_clients where active"))
    for f in lobbyists:
        if f["kind"] != "firm" or f["firm_member_ids"] or f["firm_primary_id"]:
            continue                      # an admin's (or an earlier seed's) choice stands
        labels = {_firm_phrase(_strip_tag(a)) for a in [f["name"], *(f["aliases"] or [])]}
        phrase = _firm_phrase(f["name"])
        toks = [t for t in phrase.split() if t not in _FIRM_STOP and len(t) > 1]

        squashed = phrase.replace(" ", "")

        def names_firm(text: str) -> bool:
            """"Focuspoint Communications" and "Focus Point" are one firm."""
            flat = norm_org(text).replace(" ", "")
            return bool(flat) and (squashed in flat or flat in squashed)

        def at_firm(p) -> bool:
            if names_firm(" ".join(filter(None, [p["affiliation"], p["firm"]]))):
                return True
            dom = email_domain(p["email"])
            label = dom.split(".")[0] if dom and is_private_domain(dom) else ""
            if not label:
                return False
            # A short or generic firm word ("NW") must come with the rest of
            # the name, or every @multifamilynw.org address would qualify.
            if squashed in label:
                return True
            if toks and len("".join(toks)) >= 4 and all(t in label for t in toks):
                return True
            return bool(toks) and len(toks[0]) >= 5 and toks[0] in label

        def surname_in_firm(p) -> bool:
            return bool(person_tokens(p["name"])) and person_tokens(p["name"])[-1] in toks

        scores: Counter = Counter()
        leads_2024: set[int] = set()
        for note in notes:
            firm_part, names = parse_firm_note(note)
            if _firm_phrase(firm_part) in labels:
                for n in names:
                    for p in match_person(n, persons):
                        scores[p["lobbyist_id"]] += 10      # the tracker names them for the firm
        for _, _, *labs in tracker_rows:
            if len(labs) == 2 and _firm_phrase(_strip_tag(labs[0])) in labels:
                for p in match_person(labs[1], persons):
                    scores[p["lobbyist_id"]] += 3           # paired with the firm on a donor
        for r in sheet2024:
            if not names_firm(r.get("firm") or ""):
                continue
            for p in match_person(f"{r.get('first', '')} {r.get('last', '')}", persons):
                scores[p["lobbyist_id"]] += 5               # the 2024 list's lead for the firm
                leads_2024.add(p["lobbyist_id"])
            for n in _names_in(r.get("addl_lobbyists") or ""):
                for p in match_person(n, persons):
                    scores[p["lobbyist_id"]] += 1

        members = {p["lobbyist_id"] for p in persons if at_firm(p)}
        # A firm named for people ("RAINEY & JL WILSON") keeps the people it
        # names wherever Capitol Club files them.
        members |= {pid for pid in scores if surname_in_firm(by_id[pid])}
        gone = [by_id[pid]["name"] for pid in scores if pid not in members]
        if not members:
            log.info("firm %s: no current members found (sheet names now elsewhere: %s)", f["name"], gone)
            continue
        # The 2024 list's named person for the firm is its primary (Gary
        # Oxley, not the tracker's day-to-day contact) as long as Capitol Club
        # still lists them there; the tracker's names order everyone else.
        primary = max(members, key=lambda pid: (by_id[pid]["on_capitol_club"], pid in leads_2024,
                                                scores[pid], n_clients[pid], -pid))
        dom = email_domain(by_id[primary]["email"])
        if is_private_domain(dom):
            members |= {p["lobbyist_id"] for p in persons if email_domain(p["email"]) == dom}
        ordered = [primary] + sorted(members - {primary},
                                     key=lambda pid: (-scores[pid], not surname_in_firm(by_id[pid]),
                                                      by_id[pid]["name"]))
        cur.execute("""update lobbyists set firm_primary_id = %s, firm_member_ids = %s, updated_at = now()
                       where lobbyist_id = %s and firm_primary_id is null and firm_member_ids = '{}'""",
                    (primary, ordered, f["lobbyist_id"]))
        log.info("firm %s: primary %s; %d members%s", f["name"], by_id[primary]["name"], len(ordered),
                 f" (sheet names now elsewhere: {', '.join(gone)})" if gone else "")


def _rows_raw(cur, sql, params=()):
    cur.execute(sql, params)
    return cur.fetchall()


def seed_client_leads(cur, sheet2024: list[dict]) -> None:
    """For a client several lobbyists list, the one the 2024 list names leads.

    The 2024 row's own lobbyist first, then its "Additional Lobbyists" in
    order, skipping anyone Capitol Club no longer lists for the client.
    Clients that already have a lead (an admin's choice) are left alone.
    """
    persons = [l for l in _rows(cur, "select * from lobbyists") if l["kind"] == "person"]
    listed: dict[str, set[int]] = defaultdict(set)
    by_core: dict[str, set[str]] = defaultdict(set)
    has_lead = set()
    for key, lid, lead, name in _rows_raw(cur, """select client_key, lobbyist_id, is_lead, client_name
                                                   from lobbyist_clients where active"""):
        listed[key].add(lid)
        by_core[core_org(name)].add(key)
        if lead:
            has_lead.add(key)
    set_leads = []
    for r in sheet2024:
        people = []
        for n in [f"{r.get('first', '')} {r.get('last', '')}", *_names_in(r.get("addl_lobbyists") or "")]:
            people += [p["lobbyist_id"] for p in match_person(n, persons) if p["lobbyist_id"] not in people]
        for client in [c.strip() for c in (r.get("clients") or "").split(";") if c.strip()]:
            keys = {norm_org(client)} | by_core.get(core_org(client), set())
            for key in keys:
                if key in has_lead or len(listed.get(key, ())) < 2:
                    continue
                lead = next((pid for pid in people if pid in listed[key]), None)
                if lead:
                    set_leads.append((lead, key))
                    has_lead.add(key)
    for lid, key in set_leads:
        cur.execute("""update lobbyist_clients set is_lead = true
                       where lobbyist_id = %s and client_key = %s and active""", (lid, key))
    log.info("client leads seeded from the 2024 list: %d", len(set_leads))


# ── Writing ──────────────────────────────────────────────────────────────────

def _merge(rows: list[dict], key_fields: tuple) -> dict[tuple, dict]:
    """Collapse duplicate suggestions: best score wins, evidence accumulates."""
    out: dict[tuple, dict] = {}
    for r in rows:
        k = tuple(r[f] for f in key_fields)
        if k not in out:
            out[k] = {**r, "evidence": list(r["evidence"])}
            continue
        cur = out[k]
        cur["evidence"].extend(e for e in r["evidence"] if e not in cur["evidence"])
        if r["score"] > cur["score"]:
            cur["score"], cur["method"] = r["score"], r["method"]
            if "client_name" in r:
                cur["client_name"] = r["client_name"]
    return out


def write_links(cur, table: str, key_fields: tuple, rows: dict[tuple, dict],
                extra_cols: tuple = (), status: str = "suggested", decided_by: str | None = None) -> int:
    from psycopg2.extras import execute_values
    cols = (*key_fields, *extra_cols, "method", "score", "evidence")
    values = [(*[r[c] if c != "evidence" else json.dumps(r["evidence"]) for c in cols],
               status, decided_by) for r in rows.values()]
    if not values:
        return 0
    # Human decisions are final; a machine re-run only refreshes suggestions
    # (and its own seeded decisions, marked by decided_by).
    execute_values(cur, f"""
        insert into {table} ({', '.join(cols)}, status, decided_by)
        values %s
        on conflict ({', '.join(key_fields)}) do update
          set method = excluded.method, score = excluded.score, evidence = excluded.evidence,
              status = excluded.status, decided_by = excluded.decided_by,
              decided_at = case when excluded.status = 'suggested' then null else now() end,
              updated_at = now()
          where ({table}.status = 'suggested' and {table}.method not in ('manual', 'reviewed'))
             or ({table}.decided_by = '{TRACKER_DECIDER}' and excluded.decided_by = '{TRACKER_DECIDER}')
    """, values, page_size=500)
    cur.execute(f"""update {table} set decided_at = now()
                    where status <> 'suggested' and decided_at is null""")
    return len(values)


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--refresh-pool", action="store_true", help="rebuild lobby_donor_pool first")
    ap.add_argument("--tracker", type=Path, help="Fundraising Tracker .xlsx (Lobbyist Key tab)")
    ap.add_argument("--sheet2024", type=Path, help="2024 lobby list JSON export")
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--report", type=Path, help="write all suggestions here for review")
    args = ap.parse_args()
    logging.basicConfig(level=logging.INFO, format="%(asctime)s  %(levelname)-7s  %(message)s",
                        datefmt="%H:%M:%S")

    conn = supabase_sync._connect()
    cur = conn.cursor()
    cur.execute("set statement_timeout = 0")
    if args.refresh_pool:
        log.info("donor pool refreshed: %d donors", refresh_pool(conn))

    sheet2024 = json.loads(args.sheet2024.read_text()) if args.sheet2024 else []
    # The 2024 list names which of two same-named lobbyists a label meant
    # (John Powell Jr is jcp@johnpowell.us).
    prefer_emails = {(r.get("email") or "").strip().lower() for r in sheet2024} - {""}
    if sheet2024 and not args.dry_run:
        seed_sheet2024(cur, sheet2024)
    if args.tracker and not args.dry_run:
        seed_tracker_people(cur, load_tracker(args.tracker), prefer_emails)
        seed_firm_contacts(cur, load_tracker(args.tracker), load_tracker_notes(args.tracker), sheet2024)
    if sheet2024 and not args.dry_run:
        seed_client_leads(cur, sheet2024)

    pool = _rows(cur, "select * from lobby_donor_pool")
    lobbyists = _rows(cur, "select * from lobbyists")
    lclients = _rows(cur, "select * from lobbyist_clients")
    persons = _rows(cur, "select * from committee_persons")
    log.info("pool %d donors · %d lobbyists · %d client pairs · %d committee persons",
             len(pool), len(lobbyists), len(lclients), len(persons))

    clients = {c["client_key"]: c["client_name"] for c in lclients if c["active"]}
    client_rows = match_clients(pool, clients)
    contact_links, contact_client_rows = match_committee_contacts(pool, persons, lobbyists, lclients)
    client_links = _merge(client_rows + contact_client_rows, ("donor_id", "client_key"))
    lob_links = _merge(contact_links, ("donor_id", "lobbyist_id"))

    tracker_links = {}
    if args.tracker:
        tracker_links, unresolved = tracker_suggestions(load_tracker(args.tracker), pool, lobbyists,
                                                        prefer_emails)
        for u in unresolved:
            log.warning("tracker: unresolved %s", u)

    log.info("suggestions: %d donor↔client (%s), %d committee-contact donor↔lobbyist (%s), %d tracker",
             len(client_links), dict(Counter(r["method"] for r in client_links.values())),
             len(lob_links), dict(Counter(r["method"] for r in lob_links.values())), len(tracker_links))

    if args.report:
        names = {d["donor_id"]: d["display_name"] for d in pool}
        lname = {l["lobbyist_id"]: l["name"] for l in lobbyists}
        args.report.write_text(json.dumps({
            "client_links": [{**r, "donor": names.get(r["donor_id"])} for r in client_links.values()],
            "lobbyist_links": [{**r, "donor": names.get(r["donor_id"]), "lobbyist": lname.get(r["lobbyist_id"])}
                               for r in list(lob_links.values()) + list(tracker_links.values())],
        }, indent=1, default=str))

    # A pair the tracker already decides is written once, as the tracker's,
    # carrying whatever committee-contact evidence also supports it.
    for key in set(tracker_links) & set(lob_links):
        tracker_links[key]["evidence"] += lob_links.pop(key)["evidence"]

    if args.dry_run:
        return

    n1 = write_links(cur, "donor_client_links", ("donor_id", "client_key"), client_links,
                     extra_cols=("client_name",))
    n2 = write_links(cur, "donor_lobbyist_links", ("donor_id", "lobbyist_id"), lob_links)
    n3 = write_links(cur, "donor_lobbyist_links", ("donor_id", "lobbyist_id"), tracker_links,
                     extra_cols=("is_primary",), status="confirmed", decided_by=TRACKER_DECIDER)

    # Suggestions the evidence no longer supports go away; decisions stay.
    from psycopg2.extras import execute_values
    # Pairs a person proposed (method manual/reviewed) are not the matcher's
    # to withdraw.
    mine = "status = 'suggested' and method not in ('manual', 'reviewed')"
    stale_c = [(k["donor_id"], k["client_key"]) for k in
               _rows(cur, f"select donor_id, client_key from donor_client_links where {mine}")
               if (k["donor_id"], k["client_key"]) not in client_links]
    stale_l = [(k["donor_id"], k["lobbyist_id"]) for k in
               _rows(cur, f"select donor_id, lobbyist_id from donor_lobbyist_links where {mine}")
               if (k["donor_id"], k["lobbyist_id"]) not in lob_links and
               (k["donor_id"], k["lobbyist_id"]) not in tracker_links]
    if stale_c:
        execute_values(cur, """delete from donor_client_links d using (values %s) v(donor_id, client_key)
                               where d.donor_id = v.donor_id and d.client_key = v.client_key
                                 and d.status = 'suggested'""", stale_c)
    if stale_l:
        execute_values(cur, """delete from donor_lobbyist_links d using (values %s) v(donor_id, lobbyist_id)
                               where d.donor_id = v.donor_id and d.lobbyist_id = v.lobbyist_id::bigint
                                 and d.status = 'suggested'""", stale_l)
    conn.commit()
    log.info("wrote %d client links, %d contact links, %d tracker links; removed %d + %d stale",
             n1, n2, n3, len(stale_c), len(stale_l))


def seed_sheet2024(cur, rows: list[dict]) -> None:
    """Add 2024-list lobbyists Capitol Club lacks, and their client lists.

    The list is from 2024, so its client pairs are history by default. One
    counts (active) only when the lobbyist is not on Capitol Club today AND no
    current Capitol Club lobbyist claims that client — the partner-union
    lobbyists (AFSCME, IBEW 48, OLCV) are the case this exists for.
    """
    from psycopg2.extras import execute_values
    lobbyists = _rows(cur, "select * from lobbyists")
    added, pairs = 0, {}
    for r in rows:
        first, last = (r.get("first") or "").strip(), (r.get("last") or "").strip()
        email = (r.get("email") or "").strip().lower()
        if not (first or last) or not email:
            continue
        name = f"{first} {last}".strip()
        match = [l for l in lobbyists if (l["email"] or "").lower() == email] \
            or match_person(name, lobbyists)
        if match:
            lid = match[0]["lobbyist_id"]
        else:
            cur.execute("""insert into lobbyists (name, first_name, last_name, firm, email, phone, phone_alt,
                             address, city, state, zip, source, notes)
                           values (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,'sheet_2024',%s) returning *""",
                        (name, first, last, r.get("firm") or None, email, r.get("cell") or None,
                         r.get("work") or None, r.get("address") or None, r.get("city") or None,
                         r.get("state") or None, r.get("zip") or None,
                         f"From the 2024 FuturePAC lobby list (tier {r.get('tier') or '?'}); "
                         "not on Capitol Club when imported."))
            cols = [c[0] for c in cur.description]
            row = dict(zip(cols, cur.fetchone()))
            lobbyists.append(row)
            lid = row["lobbyist_id"]
            added += 1
        for client in [c.strip() for c in re.split(r";", r.get("clients") or "") if c.strip()]:
            key = norm_org(client)
            if key:
                pairs[(lid, key)] = (lid, key, client, "sheet_2024", False)
    # The seed owns its rows: reset them, then re-apply the unclaimed rule, so
    # a client Capitol Club has since picked up stops counting for 2024 names.
    # (An admin adds a lasting pair as source 'manual' instead.)
    execute_values(cur, """insert into lobbyist_clients (lobbyist_id, client_key, client_name, source, active)
                           values %s on conflict (lobbyist_id, client_key, source)
                           do update set client_name = excluded.client_name, active = false""",
                   list(pairs.values()))
    cur.execute("""
        update lobbyist_clients s set active = true
        from lobbyists l
        where s.source = 'sheet_2024' and not s.active and l.lobbyist_id = s.lobbyist_id
          and not l.on_capitol_club
          and not exists (select 1 from lobbyist_clients c
                          where c.source = 'capitol_club' and c.active and c.client_key = s.client_key)""")
    log.info("2024 list: %d lobbyists added, %d client pairs (%d active)", added, len(pairs), cur.rowcount)


def seed_tracker_people(cur, rows, prefer_emails: set[str]) -> None:
    """Tracker labels nobody else knows become lobbyists, so the link survives.

    Firm labels ("TONKON TORP") become firm entries; person labels on neither
    Capitol Club nor the 2024 list ("Selvaggio, Mike" if he were absent)
    become person entries. A bare first name ("JEFF") is too thin to create
    anyone from and is reported instead.
    """
    lobbyists = _rows(cur, "select * from lobbyists")
    for _, _, *labels in rows:
        for label in labels:
            kind, found = resolve_label(label, lobbyists, prefer_emails)
            if found:
                continue
            raw = re.sub(r"\s+S\d\b.*$", "", label.strip())
            if kind == "firm":
                name = re.sub(r"\b(Pac West|Nw|Cfm|Jl|L&E|Communcations)\b",
                              lambda m: {"Pac West": "PAC/West", "Nw": "NW", "Cfm": "CFM", "Jl": "JL",
                                         "L&E": "L&E", "Communcations": "Communications"}[m.group(0)],
                              raw.title())
                cur.execute("""insert into lobbyists (kind, name, firm, aliases, source, notes)
                               values ('firm', %s, %s, %s, 'tracker', %s) returning *""",
                            (name, name, [raw], "Firm named in the Fundraising Tracker lobbyist key."))
            elif "," in raw:
                last, _, first = raw.partition(",")
                name = f"{first.strip().title()} {last.strip().title()}"
                cur.execute("""insert into lobbyists (kind, name, first_name, last_name, aliases, source, notes)
                               values ('person', %s, %s, %s, %s, 'tracker', %s) returning *""",
                            (name, first.strip().title(), last.strip().title(), [raw],
                             "Named in the Fundraising Tracker lobbyist key; not on Capitol Club "
                             "or the 2024 list when imported."))
            else:
                continue
            cols = [c[0] for c in cur.description]
            lobbyists.append(dict(zip(cols, cur.fetchone())))


def tracker_suggestions(rows, pool, lobbyists, prefer_emails: set[str] = frozenset()):
    by_cid = {d["committee_id"]: d for d in pool if d["committee_id"]}
    # One name can belong to several donor records (the resolver splits an
    # organization by address), and the tracker means all of them.
    by_label = defaultdict(list)
    by_core = defaultdict(list)
    for d in pool:
        for n in {d["display_name"].lower(), *(d["names"] or [])}:
            by_label[re.sub(r"\s+", " ", n).strip()].append(d)
        by_core[core_org(d["display_name"])].append(d)
    out, unresolved = {}, []
    for contributor, cid, *labels in rows:
        m = re.search(r"\((\d+)\)\s*$", contributor)
        cid = cid if cid.isdigit() else (m.group(1) if m else "")
        donors = [by_cid[cid]] if cid in by_cid else []
        donors = donors or by_label.get(re.sub(r"\s+", " ", contributor).strip().lower(), [])
        donors = donors or by_core.get(core_org(contributor), [])
        if not donors:
            unresolved.append(f"donor {contributor!r}")
            continue
        for i, label in enumerate(labels):
            _, found = resolve_label(label, lobbyists, prefer_emails)
            if not found:
                unresolved.append(f"lobbyist {label!r} (for {contributor})")
            for donor in {d["donor_id"]: d for d in donors}.values():
                for lob in found:
                    out[(donor["donor_id"], lob["lobbyist_id"])] = {
                        "donor_id": donor["donor_id"], "lobbyist_id": lob["lobbyist_id"],
                        "method": "tracker", "score": 0.97, "is_primary": i == 0,
                        "evidence": [{"type": "tracker", "contributor": contributor,
                                      "label": label, "position": i + 1}]}
    return out, unresolved


if __name__ == "__main__":
    main()
